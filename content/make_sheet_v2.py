"""Builds the second content sheet: 30 standalone (non-series) videos, one free/open-source tool each."""
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
import math

OUT = "content/freestack-30-topics-v2.xlsx"

# (title, tool, repo, replaces, keyword, script, visual, notes_for_you)
# Scripts: hook first, then only the tool and what it does. Requirements and licences live in notes_for_you.
T = [
 ("Google Photos ka free alternative", "Immich", "immich-app/immich", "Google Photos / Google One storage", "PHOTOS",
  "Google Photos ki storage full? Ab paise mat do.\nIska naam hai Immich. Google Photos ka free, open-source alternative.\nPhone ki saari photos apne khud ke computer ya hard disk pe automatic backup. Aur AI search bhi: likho beach, dog, ya birthday cake, aur woh photos turant saamne. Chehre pehchaan ke albums bhi khud bana deta hai.\nTumhari yaadein, tumhare paas. Koi monthly storage plan nahi.\nLink chahiye? Comment karo PHOTOS.",
  "Hook: phone 'Storage full' popup · Paid Google One vs Free Immich · demo: search box typing 'beach' → photo grid · GitHub card",
  "Runs on your own PC/NAS/server (Docker). AGPL-3.0."),
 ("1 minute mein apni awaaz ka AI clone", "GPT-SoVITS", "RVC-Boss/GPT-SoVITS", "ElevenLabs voice cloning", "CLONE",
  "Sirf ek minute ki awaaz, aur AI bilkul tumhari awaaz mein bolega.\nIska naam hai GPT-SoVITS. Free aur open source.\nApni ek minute ki recording do, aur ye tumhari awaaz seekh leta hai. Phir jo bhi likho, tumhari hi awaaz mein bolta hai, wahi tone, wahi andaaz. English aur kai doosri bhashayein bhi.\nApne videos ka voiceover, bina mic ke saamne baithe.\nLink chahiye? Comment karo CLONE.",
  "Hook: mic + waveform · demo: voice (Record 1 min → Clone → Bolo) · GitHub card",
  "Needs an NVIDIA GPU for training. Languages: English, Chinese, Japanese, Korean, Cantonese (no Hindi). Only clone your own or permitted voices. MIT."),
 ("Midjourney se aasaan, aur free", "Fooocus", "lllyasviel/Fooocus", "Midjourney", "IMAGE",
  "Midjourney ko bhool jao. Ye free tool usse bhi aasaan hai.\nIska naam hai Fooocus. Bilkul free aur open source.\nBas ek line likho, jaise raat mein chamakta hua Mumbai, aur ye ekdum professional AI image bana deta hai. Na koi settings, na koi jhanjhat. Realistic photos, anime, posters, sab ek click mein.\nAur ye tumhare apne computer pe chalta hai, toh jitni chaaho utni images.\nLink chahiye? Comment karo IMAGE.",
  "Hook: Midjourney renewal card · render widget (prompt → image) · GitHub card",
  "Needs an NVIDIA GPU (about 4 GB VRAM or more). GPL-3.0. Project is in maintenance mode (no new features), still works."),
 ("CapCut ka free, open-source alternative", "OpenCut", "OpenCut-app/OpenCut", "CapCut Pro", "EDIT",
  "CapCut India mein nahi chalta. Ye free editor chalta hai.\nIska naam hai OpenCut. Ek free, open-source video editor.\nTimeline, text, music, multiple tracks, sab kuch CapCut jaisa simple. Aur sabse achhi baat, koi watermark nahi, koi subscription nahi.\nTumhari videos tumhare paas rehti hain.\nReels banane waalon ke liye ekdum perfect.\nLink chahiye? Comment karo EDIT.",
  "Hook: 'App not available' screen · editor timeline mock · GitHub card",
  "Young project, still in active development; check features before promising specifics. MIT."),
 ("AI tools install karo, ek click mein", "Pinokio", "pinokiocomputer/pinokio", "Paid 1-click AI hosting apps", "CLICK",
  "AI tools install karna mushkil lagta hai? Ab sirf ek click.\nIska naam hai Pinokio. Ek free, open-source AI browser.\nIsme sainkdo AI apps ki list hai. Image generators, voice cloning, video tools, sab. Bas ek click karo, aur ye khud download, setup aur install kar deta hai. Na coding, na terminal.\nJo tools pehle sirf developers chala paate the, ab koi bhi chala sakta hai.\nLink chahiye? Comment karo CLICK.",
  "Hook: scary terminal text → one big INSTALL button · checklist · GitHub card",
  "Apps it installs still need suitable hardware (often a GPU). MIT."),
 ("Black & white photo ko rangeen banao", "DeOldify", "jantic/DeOldify", "Paid photo colorize apps", "COLOR",
  "Black and white photo, ab asli rangon mein. Bilkul free.\nIska naam hai DeOldify. Ek free, open-source AI.\nKoi bhi purani black and white photo ya video daalo, aur ye usme natural rang bhar deta hai. Aasman neela, ghaas hari, chehra asli rang ka.\nApne ghar ki purani yaadein, ya itihaas ki photos, sab rangeen.\nPehli baar dekhoge, toh rongte khade ho jayenge.\nLink chahiye? Comment karo COLOR.",
  "Hook: B&W photo → colour wipe (upscale-style slider) · GitHub card",
  "Repo is no longer actively developed but works (Colab notebooks available). MIT."),
 ("Type karna band karo, bas bolo", "Handy", "cjpais/Handy", "Wispr Flow", "TYPE",
  "Type karna band karo. Bolo, aur text khud likh jaayega.\nIska naam hai Handy. Free aur open source.\nKoi bhi app ho, email, chat ya document, bas ek shortcut dabao aur bolna shuru karo. Jo bologe, woh turant text ban jaata hai.\nAur ye poora tumhare computer pe chalta hai, bina internet.\nWispr Flow jaise paid apps ka free version.\nLink chahiye? Comment karo TYPE.",
  "Hook: hands off keyboard · wave demo (speech → typing text) · GitHub card",
  "Desktop app (Windows/macOS/Linux), offline speech-to-text. MIT."),
 ("Kisi bhi gaane se awaaz alag karo", "Ultimate Vocal Remover", "Anjok07/ultimatevocalremovergui", "Lalal.ai", "VOCAL",
  "Kisi bhi gaane se awaaz alag, music alag. Free mein.\nIska naam hai Ultimate Vocal Remover. Free aur open source.\nKoi bhi gaana daalo, aur ye singer ki awaaz aur music ko alag alag kar deta hai. Karaoke banao, remix banao, ya sirf drums aur bass nikaal lo.\nLalal jaise paid tools jo kaam karte hain, ye wahi kaam tumhare computer pe karta hai.\nMusic creators ka favourite tool.\nLink chahiye? Comment karo VOCAL.",
  "Hook: one waveform splitting into two (vocals / music) · GitHub card",
  "Faster with a GPU. Only publish remixes of songs you have rights to. MIT."),
 ("Photo ko hasaao, aankh maarwao", "LivePortrait", "KlingAIResearch/LivePortrait", "Paid face-animation apps", "LIVE",
  "Apni photo ko hasaao, aankh maarwao, sir hilwao.\nYe hai LivePortrait. Kling AI team ka free, open-source tool.\nApni ek photo do aur ek chhota sa video, aur photo bilkul usi tarah expressions karne lagti hai. Muskurana, aankh maarna, sir ghumana, sab ekdum smooth.\nMemes, cartoons aur reels ke liye mazedaar tool.\nLink chahiye? Comment karo LIVE.",
  "Hook: polaroid widget (photo starts talking) · demo checklist · GitHub card",
  "Needs a GPU. Only animate your own or permitted faces; never impersonate anyone. MIT (some model parts have their own terms)."),
 ("Purane CCTV ko smart AI camera banao", "Frigate", "blakeblackshear/frigate", "Paid cloud CCTV / smart camera plans", "CCTV",
  "Apne purane CCTV camera ko smart AI camera bana do.\nIska naam hai Frigate. Free aur open source.\nYe tumhare ghar ya dukaan ke camera ki video dekhta hai, aur pehchaanta hai ki insaan aaya, gaadi aayi, ya koi janwar. Sirf zaroori clips save karta hai, aur tumhe turant alert bhejta hai.\nSab kuch local, koi cloud subscription nahi.\nLink chahiye? Comment karo CCTV.",
  "Hook: grainy CCTV frame with AI boxes 'PERSON 97%' · checklist · GitHub card",
  "Needs an always-on PC/mini-PC and IP cameras (RTSP). MIT."),
 ("Notion AI ke paise mat do", "AppFlowy", "AppFlowy-IO/AppFlowy", "Notion AI", "NOTES",
  "Notion AI ke paise mat do. Iska free version dekho.\nIska naam hai AppFlowy. Notion ka free, open-source alternative.\nNotes, to-do lists, project boards, documents, sab ek jagah. Aur AI bhi, likhne mein madad, summary aur naye ideas.\nAur tumhara data tumhare apne device pe rehta hai.\nStudents aur teams, dono ke liye.\nLink chahiye? Comment karo NOTES.",
  "Hook: Notion AI price card · Paid vs Free · checklist · GitHub card",
  "Free AI features work with a local model; some cloud AI features are paid. AGPL-3.0."),
 ("Blurry chehra, ekdum saaf", "CodeFormer", "sczhou/CodeFormer", "Remini", "FACE",
  "Purani, dhundhli photo mein chehra saaf karne wala free AI.\nRemini jaise apps iske paise maangte hain. CodeFormer ye kaam free mein karta hai, aur iska code GitHub pe hai.\nKoi bhi purani ya blurry photo daalo, aur ye aankhein, baal, skin, sab detail ke saath wapas le aata hai.\nDaadi naani ki purani photos, ya phone ki kharab photos, sab nayi jaisi.\nLink chahiye? Comment karo FACE.",
  "Hook: blurry face → sharp (upscale slider) · GitHub card",
  "Licence is NTU S-Lab 1.0: free for non-commercial use only, so the script avoids the words 'open source'."),
 ("AI jo phone call ki tarah baat kare", "Moshi", "kyutai-labs/moshi", "Paid AI voice assistants", "TALK",
  "Ye AI tumse phone call ki tarah baat karta hai, bina ruke.\nIska naam hai Moshi. Kyutai lab ka free, open-source voice AI.\nBaaki AI pehle sunte hain, phir sochte hain, phir bolte hain. Moshi saath saath sunta aur bolta hai, bilkul insaan ki tarah. Tum beech mein bolo, woh samajh jaata hai.\nYe real-time baatcheet ka future hai.\nLink chahiye? Comment karo TALK.",
  "Hook: chat widget with two live waveforms · GitHub card",
  "Speaks English only. Needs a strong GPU or Apple Silicon Mac. Code MIT/Apache, model CC-BY."),
 ("Video ki saari khamoshi, automatic delete", "Auto-Editor", "WyattBlue/auto-editor", "Paid auto-cut tools / editor time", "CUT",
  "Video ki saari khamoshi, automatic delete.\nIska naam hai Auto-Editor. Ek free, open-source tool.\nApni lambi video daalo, aur ye jahan jahan tum chup ho, ya kuch nahi ho raha, woh saare hisse khud kaat deta hai. Result: tight, fast video, jo log end tak dekhte hain.\nAur ye Premiere, DaVinci aur Final Cut ke liye timeline bhi bana deta hai.\nLink chahiye? Comment karo CUT.",
  "Hook: long timeline with silent gaps getting snipped · GitHub card",
  "Command-line tool (no GPU needed). Public domain (Unlicense)."),
 ("Photo ki lighting badlo, bina studio", "IC-Light", "lllyasviel/IC-Light", "Studio photographer / paid relight apps", "LIGHT",
  "Normal si photo ko studio lighting mein badlo, bina studio ke.\nIska naam hai IC-Light. Ek free, open-source AI tool.\nKoi bhi photo daalo aur bolo, sunset ki roshni, neon light, ya khidki se aati dhoop. Ye photo ki poori lighting badal deta hai, aur chehra ekdum natural lagta hai.\nProduct photos, thumbnails, Instagram posts, sab ko professional look.\nLink chahiye? Comment karo LIGHT.",
  "Hook: same photo in 3 lightings side by side · GitHub card",
  "Needs a GPU (or use the Hugging Face demo). Apache-2.0."),
 ("Ye AI hasta bhi hai, gaata bhi hai", "Bark", "suno-ai/bark", "Paid expressive TTS", "BARK",
  "Ye AI sirf bolta nahi. Hasta bhi hai, gaata bhi hai.\nIska naam hai Bark. Free aur open-source text to audio model.\nLikho jo bulwana hai, aur beech mein likh do laughs ya sighs, aur ye sach mein hasega, saans lega, ruk ruk ke bolega. Gaana bhi gaa sakta hai, aur kai bhashaon mein bol sakta hai.\nRobot jaisi awaaz ka zamaana gaya.\nLink chahiye? Comment karo BARK.",
  "Hook: text with [laughs] tag → waveform · voice demo · GitHub card",
  "Faster on a GPU. Output can vary between runs. MIT."),
 ("Computer se baat karo, kaam ho jaayega", "Open Interpreter", "openinterpreter/openinterpreter", "Paid AI computer agents", "COMPUTER",
  "Apne computer se baat karo, aur woh kaam kar dega.\nIska naam hai Open Interpreter. Free aur open source.\nNormal bhasha mein bolo, jaise is folder ki saari photos ka size chhota karo, ya is Excel ka chart banao. Ye khud code likh ke tumhare computer pe woh kaam kar deta hai, aur har step chalane se pehle tumse poochta hai.\nJo kaam ghanto lagte the, ab ek line mein.\nLink chahiye? Comment karo COMPUTER.",
  "Hook: terminal widget typing a plain-English request · checklist · GitHub card",
  "Needs an AI model: API key (paid usage) or a local model. AGPL-3.0."),
 ("Apni photo ko anime bana do", "AnimeGAN", "TachibanaYoshino/AnimeGANv2", "Paid anime-filter apps", "ANIME",
  "Apni photo ko anime character bana do.\nIska naam hai AnimeGAN. Free AI tool, aur iska code GitHub pe hai.\nKoi bhi photo daalo, selfie, gali, ya pahad, aur ye use Japanese anime style mein badal deta hai. Hayao, Shinkai aur Paprika jaise styles mein.\nProfile pictures, wallpapers aur reels ke liye mazedaar.\nLink chahiye? Comment karo ANIME.",
  "Hook: selfie → anime wipe · GitHub card",
  "Licence allows non-commercial use only (commercial use needs permission), so the script avoids 'open source'."),
 ("Passport size photo, ghar baithe", "HivisionIDPhotos", "Zeyi-Lin/HivisionIDPhotos", "Photo studio / paid ID-photo apps", "PASSPORT",
  "Passport size photo ke liye studio jaana band karo.\nIska naam hai HivisionIDPhotos. Free aur open source.\nApni koi bhi normal selfie daalo, aur ye sahi size mein passport, visa ya ID card waali photo bana deta hai. Blue, white ya red background, aur print ke liye poori sheet bhi.\nForm bharne waalon ke liye bada kaam ka tool.\nLink chahiye? Comment karo PASSPORT.",
  "Hook: selfie → ID photo sheet · GitHub card",
  "Runs on CPU; also has a web demo. Check official photo rules per document. Apache-2.0."),
 ("GitHub Copilot ka free version", "Continue", "continuedev/continue", "GitHub Copilot", "COPILOT",
  "GitHub Copilot ke paise bachao. Ye free hai.\nIska naam hai Continue. Free aur open-source AI coding assistant.\nVS Code ya JetBrains mein lagao, aur ye code complete karta hai, bugs samjhata hai, aur chat mein tumhare sawaalon ke jawab deta hai, tumhare poore project ko samajh ke.\nAur tum apna pasandida AI model chun sakte ho, local model bhi.\nLink chahiye? Comment karo COPILOT.",
  "Hook: Copilot renewal card · editor widget · GitHub card",
  "Free with a local model; cloud models cost per use. Apache-2.0."),
 ("Ek photo se 1 minute ka AI video", "FramePack", "lllyasviel/FramePack", "Runway / Kling image-to-video", "ANIMATE",
  "Ek photo se poora ek minute ka AI video.\nIska naam hai FramePack. Free aur open source.\nEk image do, aur likho kya hona chahiye, jaise ladki naach rahi hai. Ye us photo se ek lamba, smooth video bana deta hai, frame by frame.\nRunway aur Kling jaise tools credits lete hain, ye tumhare computer pe chalta hai.\nLink chahiye? Comment karo ANIMATE.",
  "Hook: still photo starts moving · render widget · GitHub card",
  "Needs an NVIDIA GPU (about 6 GB VRAM or more). Apache-2.0."),
 ("Likho, aur AI awaaz bana dega", "AudioCraft", "facebookresearch/audiocraft", "Paid music/SFX libraries", "SOUND",
  "Likho baarish aur bijli, aur AI wahi awaaz bana dega.\nYe hai Meta ka AudioCraft. Free aur open source.\nIsme do kamaal ke tools hain. MusicGen likhe hue description se music banata hai, aur AudioGen sound effects. Kutte ka bhonkna, train ki awaaz, crowd ka shor, jo chahiye.\nVideos, games aur reels ke liye apni khud ki awaazein, bina stock sites pe dhoondhe.\nLink chahiye? Comment karo SOUND.",
  "Hook: text 'baarish aur bijli' → thunder flash + waveform · GitHub card",
  "Needs a GPU. Code is MIT, but model weights are CC-BY-NC (non-commercial)."),
 ("Tumhara AI second brain", "Khoj", "khoj-ai/khoj", "Paid AI note assistants", "BRAIN",
  "Tumhara apna AI second brain, jo tumhari saari notes yaad rakhe.\nIska naam hai Khoj. Free aur open source.\nApni notes, documents aur files isme daalo, aur phir sawaal pucho. Pichle mahine meeting mein kya decide hua tha? Ye tumhari hi files se jawab deta hai. Aur web search karke research bhi kar deta hai.\nTumhara personal assistant, jo kabhi kuch nahi bhoolta.\nLink chahiye? Comment karo BRAIN.",
  "Hook: chat widget asking about 'last month meeting' · GitHub card",
  "Self-host free with a local model; their hosted cloud has paid plans. AGPL-3.0."),
 ("Kharab sketch, kamaal ki painting", "Krita AI Diffusion", "Acly/krita-ai-diffusion", "Paid AI art apps", "SKETCH",
  "Tum ek kharab sa sketch banao, AI use kamaal ki painting bana dega.\nYe hai Krita AI Diffusion. Ek free, open-source plugin.\nKrita ek free drawing app hai, aur is plugin ke saath tum jo bhi draw karte ho, AI use live ek sundar artwork mein badal deta hai. Jaise jaise tum lines banate ho, painting saath saath badalti rehti hai.\nArtists ke liye ye sach mein game changer hai.\nLink chahiye? Comment karo SKETCH.",
  "Hook: stick figure → painting · GitHub card",
  "Needs a GPU (or a paid cloud option). GPL-3.0."),
 ("Itna chhota AI, itni saaf awaaz", "Kokoro", "hexgrad/kokoro", "ElevenLabs TTS", "VOICE",
  "Itna chhota AI, phir bhi itni saaf awaaz.\nIska naam hai Kokoro. Free, open-source text to speech model.\nYe sirf 82 million parameters ka hai, matlab ekdum halka aur fast. Phir bhi awaaz bilkul natural, kai voices ke saath, aur ye Hindi bhi bolta hai.\nAudiobooks, voiceovers, YouTube videos, sab ke liye.\nLink chahiye? Comment karo VOICE.",
  "Hook: tiny chip vs huge server graphic · voice demo · GitHub card",
  "Runs on CPU. Apache-2.0 (we already used it for test voiceovers)."),
 ("Photoshop ka mehenga AI feature, free", "InvokeAI", "invoke-ai/InvokeAI", "Photoshop Generative Fill / Adobe Firefly", "CANVAS",
  "Photoshop ka sabse mehenga AI feature, ab bilkul free.\nIska naam hai InvokeAI. Ek free, open-source AI studio.\nPhoto ka koi bhi hissa select karo aur likho kya chahiye. Aasman badlo, kapde badlo, ya photo ko side se badha do. AI wahi cheez naturally bana deta hai.\nLayers, canvas, brush, sab kuch ek professional editor jaisa.\nLink chahiye? Comment karo CANVAS.",
  "Hook: Adobe renewal card · Paid vs Free · canvas demo · GitHub card",
  "Needs a GPU. Apache-2.0 (models have their own licences)."),
 ("Screen pe AI dost, jo awaaz mein baat kare", "Open-LLM-VTuber", "Open-LLM-VTuber/Open-LLM-VTuber", "Paid AI companion apps", "FRIEND",
  "Tumhari screen pe ek AI dost, jo tumse awaaz mein baat karta hai.\nIska naam hai Open-LLM-VTuber. Free aur open source.\nYe ek animated character hai, jo tumhari awaaz sunta hai, sochta hai, aur bol ke jawab deta hai, chehre ke expressions ke saath. Tum beech mein bol ke use rok bhi sakte ho, bilkul asli baatcheet jaisa.\nAur ye offline bhi chal sakta hai.\nLink chahiye? Comment karo FRIEND.",
  "Hook: anime avatar with speech bubble · chat widget · GitHub card",
  "Needs a local or API AI model plus speech models; setup is technical. MIT."),
 ("Kisi bhi AI image ka prompt pata karo", "CLIP Interrogator", "pharmapsychotic/clip-interrogator", "Paid prompt packs", "PROMPT",
  "Kisi bhi AI image ka prompt pata karo.\nIska naam hai CLIP Interrogator. Free aur open source.\nKoi bhi image daalo, aur ye bata deta hai ki us image ke liye kaisa prompt likha gaya hoga. Style, lighting, camera angle, sab kuch.\nPhir wahi prompt apne AI tool mein daalo, aur waisi hi image khud banao.\nLink chahiye? Comment karo PROMPT.",
  "Hook: image → text prompt typing out · GitHub card",
  "Works on GPU or the Hugging Face demo. MIT."),
 ("Lakeer kheencho, photo badal jaayegi", "MagicQuill", "robbyant-research/MagicQuill", "Paid AI photo editors", "MAGIC",
  "Photo pe bas lakeer kheencho, AI samajh jaayega kya badalna hai.\nIska naam hai MagicQuill. Free aur open source.\nPhoto pe brush se thoda sa draw karo, jaise topi ka outline, ya koi hissa mitao, aur AI turant samajh ke wahi edit kar deta hai. Kuch jodo, kuch hatao, rang badlo, sab sirf scribble karke.\nPhoto editing itni aasaan pehle kabhi nahi thi.\nLink chahiye? Comment karo MAGIC.",
  "Hook: scribble on photo → edited result · GitHub card",
  "Needs a strong GPU. Check the repo licence before commercial use."),
 ("Painting ko ulta chalao", "Paints-UNDO", "lllyasviel/Paints-UNDO", "None (unique tool)", "ART",
  "Kisi bhi painting ko ulta chalao, aur dekho woh kaise bani thi.\nIska naam hai Paints-UNDO. Free aur open source.\nEk finished artwork do, aur ye AI dikhata hai ki use shuru se kaise banaya gaya hoga. Pehle sketch, phir lines, phir rang, step by step.\nArt seekhne waalon ke liye, aur satisfying reels ke liye, ekdum kamaal.\nLink chahiye? Comment karo ART.",
  "Hook: finished painting rewinding to a sketch · GitHub card",
  "Needs a GPU. The process it shows is AI-imagined, not the artist's real steps. Apache-2.0."),
]

wb = Workbook()
ws = wb.active
ws.title = "30 Topics"
ws.sheet_view.showGridLines = False
F = "Arial"
thin = Side(style="thin", color="D0D0D0")
box = Border(top=thin, bottom=thin, left=thin, right=thin)
fill = lambda c: PatternFill("solid", fgColor=c)
COLS = [("#", 5), ("Title", 26), ("Hook (first line)", 32), ("Free Tool", 18), ("GitHub Repo", 26),
        ("Replaces (Paid)", 20), ("Keyword", 12), ("Voice Script (Hinglish)", 70), ("Visual Notes", 30),
        ("Notes for you (NOT in script)", 34), ("Status", 15)]
for i, (_, w) in enumerate(COLS, 1):
    ws.column_dimensions[chr(64 + i)].width = w
last = chr(64 + len(COLS))
ws.merge_cells(f"A1:{last}1")
ws["A1"] = "FREESTACK AI  ·  30 STANDALONE VIDEO TOPICS (V2)"
ws["A1"].font = Font(name=F, size=18, bold=True, color="FFE500")
ws["A1"].fill = fill("111111")
ws["A1"].alignment = Alignment(vertical="center", indent=1)
ws.row_dimensions[1].height = 36
ws.merge_cells(f"A2:{last}2")
ws["A2"] = ("One free / open-source tool per video. Scripts = hook + tool + what it does + CTA. "
            "Requirements and licences are in the last notes column only. Repos checked on GitHub, 16 Oct 2026.")
ws["A2"].font = Font(name=F, size=10, italic=True, color="555555")
ws["A2"].alignment = Alignment(vertical="center", indent=1)
HDR, FIRST = 4, 5
for i, (h, _) in enumerate(COLS, 1):
    c = ws.cell(row=HDR, column=i, value=h)
    c.font = Font(name=F, bold=True, color="FFE500", size=10)
    c.fill = fill("111111")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = box
ws.row_dimensions[HDR].height = 30

def lines(text, width):
    return sum(max(1, math.ceil(len(p) / (width * 1.25))) for p in str(text).split("\n"))

for n, (title, tool, repo, repl, kw, script, visual, notes) in enumerate(T, 1):
    r = HDR + n
    hook = script.split("\n")[0]
    vals = [n, title, hook, tool, repo, repl, kw, script, visual, notes, "To do"]
    bg = "FFFFFF" if n % 2 else "F7F7F5"
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = Font(name=F, size=10, bold=(i in (1, 2, 7)), color=("FF3D7F" if i == 7 else "111111"))
        c.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center" if i in (1, 7, 11) else "left")
        c.fill = fill(bg)
        c.border = box
    ws.cell(row=r, column=1).font = Font(name=F, size=14, bold=True)
    ws.cell(row=r, column=10).font = Font(name=F, size=9, italic=True, color="7A4B00")
    ws.row_dimensions[r].height = min(409, max(lines(t, w) for t, w in [(script, 70), (visual, 30), (notes, 34), (hook, 32)]) * 13 + 8)
END = HDR + len(T)

statuses = [("To do", "E7E6E6"), ("Script ready", "FFF2A8"), ("Voice recorded", "CFE3FF"), ("Video built", "E5D4FF"), ("Posted", "C9F2D0")]
dv = DataValidation(type="list", formula1='"' + ",".join(s for s, _ in statuses) + '"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"K{FIRST}:K{END}")
for name, col in statuses:
    ws.conditional_formatting.add(f"K{FIRST}:K{END}", CellIsRule(operator="equal", formula=[f'"{name}"'], fill=fill(col), font=Font(name=F, bold=True)))
ws.freeze_panes = f"C{FIRST}"
ws.auto_filter.ref = f"A{HDR}:{last}{END}"
nr = END + 2
for k, v in [("HOW TO USE", "Change only the Status column. Copy the Voice Script into your voice tool, send me the audio, I build the reel."),
             ("Rule", "Scripts never mention requirements (GPU, API key, setup). Check the notes column before posting so nothing in the video is wrong."),
             ("Licences", "CodeFormer and AnimeGAN are free for non-commercial use only, so their scripts say 'free' but not 'open source'.")]:
    ws.merge_cells(start_row=nr, start_column=1, end_row=nr, end_column=2)
    ws.cell(row=nr, column=1, value=k).font = Font(name=F, size=10, bold=True)
    ws.merge_cells(start_row=nr, start_column=3, end_row=nr, end_column=8)
    ws.cell(row=nr, column=3, value=v).font = Font(name=F, size=10)
    nr += 1
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = f"{HDR}:{HDR}"
wb.save(OUT)
print("saved", OUT, len(T), "topics")
