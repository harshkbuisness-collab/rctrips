"""Builds the 30-day plan v3: tabs for Reels, Carousels and Posts (one of each per day)."""
import math
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from plan_v3_data import REELS, CAROUSELS, POSTS

OUT = "content/freestack-30day-plan-v3.xlsx"
F = "Arial"
thin = Side(style="thin", color="D0D0D0")
box = Border(top=thin, bottom=thin, left=thin, right=thin)
fill = lambda c: PatternFill("solid", fgColor=c)
STAT = [("To do", "E7E6E6"), ("Script ready", "FFF2A8"), ("Voice recorded", "CFE3FF"), ("Built", "E5D4FF"), ("Posted", "C9F2D0")]


def sheet(ws, title, sub, cols, rows, center=(), bold=(), note_col=None):
    ws.sheet_view.showGridLines = False
    last = chr(64 + len(cols))
    for i, (_, w) in enumerate(cols, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.merge_cells(f"A1:{last}1"); ws["A1"] = title
    ws["A1"].font = Font(name=F, size=16, bold=True, color="FFE500"); ws["A1"].fill = fill("111111")
    ws["A1"].alignment = Alignment(vertical="center", indent=1); ws.row_dimensions[1].height = 32
    ws.merge_cells(f"A2:{last}2"); ws["A2"] = sub
    ws["A2"].font = Font(name=F, size=10, italic=True, color="555555")
    for i, (h, _) in enumerate(cols, 1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(name=F, bold=True, color="FFE500", size=10); c.fill = fill("111111")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = box
    ws.row_dimensions[4].height = 28
    for n, vals in enumerate(rows, 1):
        r = 4 + n
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = Font(name=F, size=10, bold=(i in bold))
            c.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center" if i in center else "left")
            c.fill = fill("FFFFFF" if n % 2 else "F7F7F5"); c.border = box
        ws.cell(row=r, column=1).font = Font(name=F, size=13, bold=True)
        if note_col:
            ws.cell(row=r, column=note_col).font = Font(name=F, size=9, italic=True, color="7A4B00")
        ws.row_dimensions[r].height = min(409, max(
            sum(max(1, math.ceil(len(p) / (cols[i][1] * 1.2))) for p in str(v).split("\n"))
            for i, v in enumerate(vals)) * 13 + 8)
    end = 4 + len(rows)
    dv = DataValidation(type="list", formula1='"' + ",".join(s for s, _ in STAT) + '"')
    ws.add_data_validation(dv); dv.add(f"{last}5:{last}{end}")
    for s, col in STAT:
        ws.conditional_formatting.add(f"{last}5:{last}{end}", CellIsRule(operator="equal", formula=[f'"{s}"'], fill=fill(col), font=Font(name=F, bold=True)))
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{last}{end}"


wb = Workbook()
ws = wb.active; ws.title = "Reels"
rows = []
for d, (typ, title, tool, repo, kw, script, notes) in enumerate(REELS, 1):
    words = len(script.split())
    rows.append([d, typ, title, script.split("\n")[0], tool, repo, kw, script, f"{words} words ≈ {round(words / 2.6)} s", notes, "To do"])
sheet(ws, "REELS · 30 DAYS · under 20 s, friend tone", "20 free-tool Reels (reach) + 10 Claude showcase Reels (buyers). Notes column is for you, never in the script.",
      [("Day", 6), ("Type", 9), ("Title", 26), ("Hook", 30), ("Tool", 16), ("Repo", 24), ("Keyword", 11), ("Voice Script", 60), ("Length", 12), ("Notes (not in script)", 30), ("Status", 13)],
      rows, center=(1, 2, 7, 9), bold=(3, 7), note_col=10)

ws = wb.create_sheet("Carousels")
rows = []
for d, (topic, cover, slides, cta, kw) in enumerate(CAROUSELS, 1):
    rows.append([d, topic, cover, "\n".join(f"{i}. {s}" for i, s in enumerate(slides, 1)), len(slides), cta, kw, "To do"])
sheet(ws, "CAROUSELS · 30 DAYS · notebook explainers", "Explainers, analogies, frameworks and honest guides (save-worthy). Style: graph-paper notebook, FreeStack brand, mascot stamp.",
      [("Day", 6), ("Topic", 24), ("Cover title (hook)", 32), ("Slides outline", 60), ("Slides", 8), ("CTA", 14), ("Keyword", 11), ("Status", 13)],
      rows, center=(1, 5, 7), bold=(3, 7))

ws = wb.create_sheet("Posts")
rows = [[d, fmt, head, body, cap, "To do"] for d, (fmt, head, body, cap) in enumerate(POSTS, 1)]
sheet(ws, "POSTS · 30 DAYS · single image", "One quick, shareable image a day: paid-vs-free cards, myths, prompts, polls, before/after, behind the scenes.",
      [("Day", 6), ("Format", 18), ("Headline (on image)", 34), ("Body (on image)", 40), ("Caption", 44), ("Status", 13)],
      rows, center=(1,), bold=(3,))

wb.save(OUT)
print("saved", OUT, len(REELS), len(CAROUSELS), len(POSTS))
